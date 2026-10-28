# -*- coding: utf-8 -*-
"""批量导入：把「角色 + 时点 + 事项 + 提醒内容」的计划表，一键导成 notice 提醒模板。

⚠️ 2026-10-26 项目换轨：本脚本跑在**新主线** `app/`（StateStore）上，
   旧的 `src/imtask` 已归档到 `_legacy/`，⛔ 不要再往那边写。

幂等设计（可以反复跑）
----------------------
    角色  按名称唯一（已存在则复用，⛔ 不重复建）
    成员  按 (角色, 手机号) 唯一（已存在则跳过）
    模板  按**模板名**唯一（同名则更新文案与时刻，不新建）

一表一时点：本项目一个模板只承载一个 `text_tpl` + 1~3 个 `notice_times`，
春促表 5 个时点文案各不相同 → 建 5 个模板（每个 1 个时刻）。

用法
----
    python scripts/import_plan.py --file "C:\\path\\春促角色清单.xlsx" \\
        --phone 13900000000 --name 张三 [--dry-run] [--preview]
"""

from __future__ import annotations

import argparse
import sys
from datetime import time
from pathlib import Path
from typing import Any, Dict, List

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from app.config import Config  # noqa: E402
from app.scheduler import render_text_tpl  # noqa: E402
from app.state import BizError, StateStore  # noqa: E402

COLUMN_ALIASES = {
    "role": ("角色",),
    "time": ("时间", "时点", "提醒时间"),
    "item": ("事项", "事项名称"),
    "text": ("提醒内容", "内容", "文案"),
}
ACTOR = "import"

#: 2026-10-26 业务方定档：文案里的「一、文字提醒：」「二、数据提醒：」两个分段标签**不进消息**，
#: 只保留标签下面的正文；两段之间留一个空行分隔（否则两串编号会连在一起分不清）。
_SECTION_LABELS = ("一、文字提醒", "二、数据提醒", "三、数据提醒")


def clean_text(raw: str, keep_labels: bool = False) -> str:
    """去掉分段标签行；保留原始换行与标点，⛔ 不改任何正文。"""
    if keep_labels:
        return raw
    out: List[str] = []
    for line in raw.split("\n"):
        stripped = line.strip()
        is_label = stripped.rstrip("：:") in _SECTION_LABELS
        if is_label:
            if out and out[-1].strip():
                out.append("")
            continue
        out.append(line.rstrip())
    # 折叠连续空行、去掉首尾空行
    cleaned: List[str] = []
    for line in out:
        if not line and cleaned and not cleaned[-1]:
            continue
        cleaned.append(line)
    while cleaned and not cleaned[0]:
        cleaned.pop(0)
    while cleaned and not cleaned[-1]:
        cleaned.pop()
    return "\n".join(cleaned)


def _fmt_time(value: Any) -> str:
    if isinstance(value, time):
        return value.strftime("%H:%M")
    text = "" if value is None else str(value).strip().replace("：", ":")
    if not text:
        return ""
    if "点" in text:
        hour, _, minute = text.partition("点")
        return "%02d:%02d" % (int(hour), int(minute.replace("分", "") or 0))
    parts = text.split(":")
    if len(parts) >= 2:
        return "%02d:%02d" % (int(parts[0]), int(parts[1]))
    raise ValueError("无法识别的时间：%r" % (value,))


def load_rows(path: Path) -> Dict[str, Any]:
    from openpyxl import load_workbook

    workbook = load_workbook(path, data_only=True)
    sheet = workbook.active
    headers = [sheet.cell(row=1, column=i).value
               for i in range(1, sheet.max_column + 1)]
    index: Dict[str, int] = {}
    for key, aliases in COLUMN_ALIASES.items():
        for position, name in enumerate(headers, start=1):
            if name is not None and str(name).strip() in aliases:
                index[key] = position
                break
    for key in ("time", "text"):
        if key not in index:
            raise SystemExit("表格缺少必需列：%s（表头：%s）" % (key, headers))

    rows: List[Dict[str, str]] = []
    last_role = ""
    for row_no in range(2, sheet.max_row + 1):
        def cell(key: str):
            return (sheet.cell(row=row_no, column=index[key]).value
                    if key in index else None)

        text = "" if cell("text") is None else str(cell("text"))
        if not text.strip():
            continue
        role = "" if cell("role") is None else str(cell("role")).strip()
        if role:
            last_role = role
        rows.append({"role": role or last_role, "time": _fmt_time(cell("time")),
                     "item": "" if cell("item") is None
                             else str(cell("item")).strip(),
                     "text": text})
    roles = sorted({row["role"] for row in rows if row["role"]})
    return {"headers": headers, "index": index, "roles": roles, "rows": rows}


def template_name(role: str, item: str, hhmm: str) -> str:
    return "%s · %s %s" % (role, hhmm, item or "提醒")


def ensure_role(state: StateStore, name: str, dry: bool) -> str:
    for row in state.list_roles():
        if row["name"] == name:
            print("  角色已存在：%s（%s，%d 人）"
                  % (name, row["role_id"], row["member_count"]))
            return row["role_id"]
    if dry:
        print("  角色将新建：%s" % name)
        return "(dry-run)"
    created = state.create_role(name, "由春促计划表导入", ACTOR)
    print("  角色已新建：%s（%s）" % (name, created["role_id"]))
    return created["role_id"]


def ensure_member(state: StateStore, role_id: str, phone: str, name: str,
                  dry: bool) -> None:
    if dry or role_id == "(dry-run)":
        print("  成员将加入：%s / %s（%s）" % (role_id, phone, name))
        return
    try:
        state.add_role_member(role_id, phone, name, None, "import", ACTOR)
        print("  成员已加入：%s（%s）" % (phone, name))
    except BizError as exc:
        if exc.status_code == 409:
            print("  成员已存在，跳过：%s（%s）" % (phone, name))
        else:
            raise


def upsert_template(state: StateStore, *, name: str, role_id: str,
                    hhmm: str, text: str, dry: bool) -> str:
    payload = {
        "name": name, "kind": "notice", "role_ids": [role_id],
        "text_tpl": text, "period": "day", "notice_times": [hhmm],
        "enabled": True,
    }
    existing = None
    for item in state.list_templates(q=name, size=100)["items"]:
        if item["name"] == name:
            existing = item
            break
    if existing:
        if dry:
            return "将更新"
        state.update_template(existing["template_id"], payload, ACTOR)
        return "已更新"
    if dry:
        return "将新建"
    state.create_template(payload, ACTOR)
    return "已新建"


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--file", required=True)
    parser.add_argument("--phone", action="append", default=[])
    parser.add_argument("--name", action="append", default=[])
    parser.add_argument("--dry-run", action="store_true")
    parser.add_argument("--preview", action="store_true",
                        help="导入后打印「用户实际收到的一条消息」")
    parser.add_argument("--keep-labels", action="store_true",
                        help="保留「一、文字提醒：」这类分段标签（默认去掉）")
    args = parser.parse_args()

    data = load_rows(Path(args.file))
    print("表头：%s" % data["headers"])
    print("识别列：%s" % data["index"])
    print("角色：%s ｜ 待导入 %d 条\n" % ("、".join(data["roles"]),
                                        len(data["rows"])))

    cfg = Config.from_env()
    state = StateStore(cfg.db_path)
    print("库：%s\n" % cfg.db_path)

    print("① 角色")
    role_ids = {name: ensure_role(state, name, args.dry_run)
                for name in data["roles"]}

    print("\n② 成员")
    names = args.name + [""] * (len(args.phone) - len(args.name))
    for role in data["roles"]:
        for phone, person in zip(args.phone, names):
            ensure_member(state, role_ids[role], phone, person, args.dry_run)

    print("\n③ 提醒模板（一个时点一个）")
    for row in data["rows"]:
        name = template_name(row["role"], row["item"], row["time"])
        text = clean_text(row["text"], args.keep_labels)
        action = upsert_template(state, name=name,
                                 role_id=role_ids.get(row["role"], ""),
                                 hhmm=row["time"], text=text,
                                 dry=args.dry_run)
        print("  %-6s %-44s %s  正文 %d 字（原 %d 字）"
              % (action, name, row["time"], len(text), len(row["text"])))

    if args.preview and not args.dry_run:
        print("\n④ 预览：用户实际收到的消息（notice 不加任何前缀）")
        for item in state.list_templates(size=100)["items"]:
            if not item["name"].startswith("片长/分支局长"):
                continue
            text = render_text_tpl(
                item["text_tpl"], member_name="张三", member_dept=None,
                template_name=item["name"], deadline_at=None,
                period=item.get("period") or "day")
            print("\n--- %s ---\n%s" % (item["name"], text))

    print("\n完成%s" % ("（--dry-run，未写库）" if args.dry_run else ""))


if __name__ == "__main__":
    main()
